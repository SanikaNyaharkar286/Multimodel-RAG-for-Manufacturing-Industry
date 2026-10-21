from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

root = Path(__file__).parent
source = root / "BACKEND_AND_AGENT_GUIDE.md"
output = root / "Manufacturing_Equipment_Copilot_Backend_Guide.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = None

def add_code(lines):
    for line in lines:
        p = doc.add_paragraph()
        p.style = "No Spacing"
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)

lines = source.read_text(encoding="utf-8").splitlines()
i = 0
first = True
while i < len(lines):
    line = lines[i]
    if line.startswith("```"):
        i += 1
        block = []
        while i < len(lines) and not lines[i].startswith("```"):
            block.append(lines[i])
            i += 1
        add_code(block)
    elif line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line[2:].strip())
        first = False
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.startswith("- "):
        doc.add_paragraph(line[2:].strip(), style="List Bullet")
    elif line and not line.startswith("|"):
        doc.add_paragraph(line)
    i += 1

doc.core_properties.title = "Manufacturing Equipment Copilot Backend Guide"
doc.core_properties.subject = "Backend, agent, retrieval, and ingestion documentation"
doc.save(output)
print(output)
